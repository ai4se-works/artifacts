#!/usr/bin/env python3
"""Extract anonymized derived CSVs from internal handoff xlsx (not shipped).

Reads (relative to repo root):
  99-archive/handoff/03-evidence/maturity/ai4se-maturity-before-after-compare-visualization.xlsx
  99-archive/handoff/03-evidence/metrics/ai4se-effectiveness-measurement-summary.xlsx

Writes next to this pack under derived:
  maturity_domain_scores.csv
  maturity_capabilities.csv
  effectiveness_pair_summary.csv
  effectiveness_aggregate.csv

Never copies the 成员 (member name) column — pairs become Pair-1..Pair-6.
"""

from __future__ import annotations

import csv
import math
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

REPO_ROOT = Path(__file__).resolve().parents[5]
MATURITY_XLSX = (
    REPO_ROOT
    / "99-archive/handoff/03-evidence/maturity"
    / "ai4se-maturity-before-after-compare-visualization.xlsx"
)
EFFECTIVENESS_XLSX = (
    REPO_ROOT
    / "99-archive/handoff/03-evidence/metrics"
    / "ai4se-effectiveness-measurement-summary.xlsx"
)
OUT_DIR = Path(__file__).resolve().parents[1] / "derived"
    REPO_ROOT / "docs/research/fse-icse-2027/public-data/derived"
)

# English labels aligned to Chinese 能力域 / 能力项 in the maturity workbook.
DOMAIN_NAME_EN = {
    "D1": "Efficiency Foundation",
    "D2": "Process Layer",
    "D3": "Engineering Methods & Discipline",
    "D4": "Tooling Layer",
    "D5": "Human–Agent Collaboration",
    "D6": "Organization & Culture",
}

CAPABILITY_NAME_EN = {
    "D1.1": "Value Goals & Outcome Hypotheses",
    "D1.2": "Quality & Trust Baseline",
    "D1.3": "Efficiency & Flow Measurement",
    "D2.1": "AI4SE Workflow Design",
    "D2.2": "Small-Batch Task Decomposition",
    "D2.3": "Evidence-Driven Review/Ship Gates",
    "D3.1": "Spec-Driven Expression",
    "D3.2": "Context Organization",
    "D3.3": "Verification & Evidence Discipline",
    "D4.1": "Agent Runtime & Permissions",
    "D4.2": "AI-Accessible Engineering Context",
    "D4.3": "Automated Verification & Toolchain Integration",
    "D5.1": "Collaboration Duties & Independent Verification",
    "D5.2": "HITL/HOTL/HOOL Risk Routing",
    "D5.3": "Accountability & Audit Trails",
    "D6.1": "Practice Champions & Training",
    "D6.2": "Community & Knowledge Assets",
    "D6.3": "Adoption Confidence & Behavior Change",
}

LEVEL_RE = re.compile(r"^(L[1-5])")


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    return float(text)


def _fmt(value: float | None, digits: int = 6) -> str:
    if value is None:
        return ""
    if abs(value - round(value)) < 1e-12:
        return str(int(round(value)))
    text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
    return text


def _level_code(label) -> str:
    if label is None:
        return ""
    match = LEVEL_RE.match(str(label).strip())
    return match.group(1) if match else str(label).strip()


def _find_header_row(ws, required_token: str) -> int:
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if required_token in cells:
            return idx
    raise ValueError(f"Header containing {required_token!r} not found in {ws.title}")


def extract_maturity_domains(wb) -> list[dict]:
    ws = wb["Domain Scores"]
    header_row = _find_header_row(ws, "域")
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        domain_id = row[0]
        if not domain_id or not str(domain_id).startswith("D"):
            continue
        domain_id = str(domain_id).strip()
        before = _num(row[2])
        after = _num(row[3])
        delta = _num(row[4])
        before_l1 = row[5]
        after_l1 = row[6]
        before_l3p = row[7]
        after_l3p = row[8]
        notes = (
            f"L1 capabilities {before_l1}->{after_l1}; "
            f"L3+ capabilities {before_l3p}->{after_l3p}"
        )
        rows.append(
            {
                "domain_id": domain_id,
                "domain_name_en": DOMAIN_NAME_EN[domain_id],
                "before": _fmt(before, 2),
                "after": _fmt(after, 2),
                "delta": _fmt(delta, 2),
                "signal_type": "process_signal",
                "notes": notes,
            }
        )
    return rows


def extract_maturity_capabilities(wb) -> list[dict]:
    ws = wb["Capability Heatmap"]
    header_row = _find_header_row(ws, "能力项ID")
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        domain_id = row[0]
        cap_id = row[2]
        if not domain_id or not cap_id:
            continue
        domain_id = str(domain_id).strip()
        cap_id = str(cap_id).strip()
        before = _num(row[4])
        after = _num(row[5])
        delta = _num(row[6])
        rows.append(
            {
                "domain_id": domain_id,
                "capability_id": cap_id,
                "capability_name_en": CAPABILITY_NAME_EN[cap_id],
                "before": _fmt(before, 0),
                "after": _fmt(after, 0),
                "delta": _fmt(delta, 0),
                "before_level": _level_code(row[7]),
                "after_level": _level_code(row[8]),
                "signal_type": "process_signal",
            }
        )
    return rows


def _pair_id_from_cell(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    match = re.fullmatch(r"Pair\((\d+)\)", text)
    if not match:
        return None
    return f"Pair-{int(match.group(1))}"


def extract_effectiveness_pairs(wb) -> list[dict]:
    ws = wb["汇总对比报告"]
    header_row = _find_header_row(ws, "Pair")
    header = [c for c in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    # Source columns (row 5): Pair, 成员, TP Before, TP After, Δ提效 TP,
    # CT1 P75 Before, CT1 P75 After, Δ提效 CT1 P75, 速度倍数 CT1, 缩短% CT1,
    # CT2 P75 Before, CT2 P75 After, Δ提效 CT2 P75, 速度倍数 CT2
    # Note: no CT2 缩短% column in xlsx — derived as (before-after)/before.
    assert header[1] == "成员", "expected 成员 column at index 1 (skipped for anonymity)"
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        pair_id = _pair_id_from_cell(row[0])
        if not pair_id:
            break
        # Intentionally ignore row[1] (成员 names).
        tp_before = _num(row[2])
        tp_after = _num(row[3])
        tp_delta = _num(row[4])
        ct1_b = _num(row[5])
        ct1_a = _num(row[6])
        ct1_speedup = _num(row[8])
        ct1_reduction = _num(row[9])
        ct2_b = _num(row[10])
        ct2_a = _num(row[11])
        ct2_speedup = _num(row[13])
        ct2_reduction = None
        if ct2_b is not None and ct2_a is not None and ct2_b != 0:
            ct2_reduction = (ct2_b - ct2_a) / ct2_b
        rows.append(
            {
                "pair_id": pair_id,
                "tp_before": _fmt(tp_before, 0),
                "tp_after": _fmt(tp_after, 0),
                "tp_delta_improve": _fmt(tp_delta, 6),
                "ct1_p75_before": _fmt(ct1_b, 4),
                "ct1_p75_after": _fmt(ct1_a, 4),
                "ct1_speedup": _fmt(ct1_speedup, 6),
                "ct1_reduction_pct": _fmt(ct1_reduction, 6),
                "ct2_p75_before": _fmt(ct2_b, 4),
                "ct2_p75_after": _fmt(ct2_a, 4),
                "ct2_speedup": _fmt(ct2_speedup, 6),
                "ct2_reduction_pct": _fmt(ct2_reduction, 6),
                "signal_type": "expert_estimate",
            }
        )
    return rows


def extract_effectiveness_aggregate(wb) -> list[dict]:
    """Cross-pair median / IQR of delta-improve from section 二 of 汇总对比报告."""
    ws = wb["汇总对比报告"]
    # Locate aggregate table header
    header_row = None
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "指标" and row[1] == "Δ提效 中位数":
            header_row = idx
            break
    if header_row is None:
        raise ValueError("Aggregate header row not found")

    metric_map = {
        "Throughput（ΣSP 产能提效）": "tp_delta_improve",
        "CT1 → last commit（P75，速度提效，主报）": "ct1_speed_delta_improve_p75",
        "CT2 → QA test done（P75，速度提效，主报）": "ct2_speed_delta_improve_p75",
    }
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        label = row[0]
        if not label or label not in metric_map:
            # Stop after primary P75 block (skip P50 annex rows).
            if label and str(label).startswith("CT1 → last commit（P50"):
                break
            continue
        metric = metric_map[str(label)]
        median = _num(row[1])
        p25 = _num(row[2])
        p75 = _num(row[3])
        iqr = _num(row[4])
        vmin = _num(row[5])
        vmax = _num(row[6])
        for stat, value in (
            ("median", median),
            ("p25", p25),
            ("p75", p75),
            ("iqr", iqr),
            ("min", vmin),
            ("max", vmax),
        ):
            rows.append(
                {
                    "metric": metric,
                    "stat": stat,
                    "value": _fmt(value, 6),
                    "signal_type": "expert_estimate",
                }
            )
    return rows


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main() -> int:
    if not MATURITY_XLSX.is_file():
        print(f"Missing maturity xlsx: {MATURITY_XLSX}", file=sys.stderr)
        return 1
    if not EFFECTIVENESS_XLSX.is_file():
        print(f"Missing effectiveness xlsx: {EFFECTIVENESS_XLSX}", file=sys.stderr)
        return 1

    maturity_wb = load_workbook(MATURITY_XLSX, data_only=True)
    effect_wb = load_workbook(EFFECTIVENESS_XLSX, data_only=True)

    domain_rows = extract_maturity_domains(maturity_wb)
    cap_rows = extract_maturity_capabilities(maturity_wb)
    pair_rows = extract_effectiveness_pairs(effect_wb)
    agg_rows = extract_effectiveness_aggregate(effect_wb)

    n_domain = write_csv(
        OUT_DIR / "maturity_domain_scores.csv",
        [
            "domain_id",
            "domain_name_en",
            "before",
            "after",
            "delta",
            "signal_type",
            "notes",
        ],
        domain_rows,
    )
    n_cap = write_csv(
        OUT_DIR / "maturity_capabilities.csv",
        [
            "domain_id",
            "capability_id",
            "capability_name_en",
            "before",
            "after",
            "delta",
            "before_level",
            "after_level",
            "signal_type",
        ],
        cap_rows,
    )
    n_pair = write_csv(
        OUT_DIR / "effectiveness_pair_summary.csv",
        [
            "pair_id",
            "tp_before",
            "tp_after",
            "tp_delta_improve",
            "ct1_p75_before",
            "ct1_p75_after",
            "ct1_speedup",
            "ct1_reduction_pct",
            "ct2_p75_before",
            "ct2_p75_after",
            "ct2_speedup",
            "ct2_reduction_pct",
            "signal_type",
        ],
        pair_rows,
    )
    n_agg = write_csv(
        OUT_DIR / "effectiveness_aggregate.csv",
        ["metric", "stat", "value", "signal_type"],
        agg_rows,
    )

    print(f"maturity_domain_scores.csv: {n_domain} data rows")
    print(f"maturity_capabilities.csv: {n_cap} data rows")
    print(f"effectiveness_pair_summary.csv: {n_pair} data rows")
    print(f"effectiveness_aggregate.csv: {n_agg} data rows")
    print(f"wrote under {OUT_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
